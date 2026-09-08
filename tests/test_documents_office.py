from dataclasses import replace
from datetime import datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import subprocess
import threading
from zipfile import ZipFile, ZIP_DEFLATED

import fitz
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from sandevistan_read import documents, retrieval, services
from sandevistan_read.database import Database, json_load


def test_chinese_question_keeps_late_technical_search_terms(monkeypatch):
    rows = [{'id': f'c{i}', 'source_id': 'source', 'content': '无关内容',
             'locator_json': '{}', 'embedding_json': '[1, 0]'} for i in range(8)]
    rows[-1]['content'] = 'Offline-first ASR and TTS for Linux and Windows.'
    monkeypatch.setattr(retrieval, 'DB', SimpleNamespace(fetchall=lambda *args: rows))
    monkeypatch.setattr(retrieval.EMBEDDINGS, 'encode', lambda *args, **kwargs: [[1, 0]])
    result = retrieval.retrieve('notebook', '请详细说明这个本地语音工具提供哪些功能以及支持的操作系统，尤其是 ASR 和 TTS。', ['source'])
    assert result[0]['id'] == 'c7'


@pytest.fixture
def local_paths(tmp_path, monkeypatch):
    paths = replace(documents.PATHS, root=tmp_path, renders=tmp_path / 'renders', libreoffice_profiles=tmp_path / 'profiles')
    monkeypatch.setattr(documents, 'PATHS', paths)
    monkeypatch.setattr(services, 'PATHS', paths)
    monkeypatch.setattr(documents, '_convert_office_to_pdf', lambda *args: None)
    return paths


def test_xlsx_values_formulas_sheets_and_numeric_retrieval(local_paths):
    book = Workbook()
    sheet = book.active
    sheet.title = '资产'
    sheet.append(['日期', '收益率', '数值', '合计'])
    sheet.append([datetime(2026, 9, 5), .1234, 42, '=C2*2'])
    sheet['B2'].number_format = '0.00%'
    sheet['D3'] = '=C2*3'
    sheet.merge_cells('A4:B4')
    sheet['A4'] = '合并标签'
    book.create_sheet('其他').append(['标签', 123])
    path = local_paths.root / 'sample.xlsx'
    book.save(path)
    # Supply an Excel-style saved formula value without evaluating a formula.
    with ZipFile(path) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    contents['xl/worksheets/sheet1.xml'] = contents['xl/worksheets/sheet1.xml'].replace(b'<f>C2*2</f><v></v>', b'<f>C2*2</f><v>84</v>')
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, value in contents.items():
            archive.writestr(name, value)
    parsed = documents.parse_document(path, 'xlsx')
    text = '\n'.join(block.text for block in parsed.blocks)
    assert parsed.page_count == 2
    assert '2026-09-05' in text and '12.34%' in text and 'D2: 84' in text
    assert '公式未缓存：=C2*3' in text
    assert parsed.metadata['merged_cells']['资产'] == ['A4:B4']
    assert parsed.metadata['warnings'][0]['count'] == 1
    assert all(block.locator['cell_range'] for block in parsed.blocks)
    assert all(retrieval.is_quality_chunk({'content': block.text, 'locator': block.locator}) for block in parsed.blocks)


@pytest.mark.parametrize('failure', [FileNotFoundError('missing'), subprocess.TimeoutExpired('soffice', 180), UnicodeDecodeError('utf-8', b'\xff', 0, 1, 'bad encoding')])
def test_docx_keeps_body_order_when_preview_fails(local_paths, monkeypatch, failure):
    doc = Document()
    doc.add_heading('第一节', 1)
    doc.add_paragraph('表格之前。')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '指标'
    table.cell(0, 1).text = '数值'
    table.cell(1, 0).text = '收益'
    table.cell(1, 1).text = '42'
    doc.add_heading('第二节', 1)
    doc.add_paragraph('表格之后。')
    path = local_paths.root / 'sample.docx'
    doc.save(path)
    def broken(*args):
        raise failure
    monkeypatch.setattr(documents, '_convert_office_to_pdf', broken)
    parsed = documents.parse_document(path, 'docx')
    text = '\n'.join(block.text for block in parsed.blocks)
    assert text.index('表格之前') < text.index('指标') < text.index('第二节')
    assert next(block for block in parsed.blocks if block.locator['kind'] == 'table').locator['section'] == '第一节'
    assert parsed.metadata['warnings'][0]['code'] == 'office_preview_unavailable'
    assert documents.chunk_blocks(parsed.blocks)


def test_pptx_grouped_text_and_picture_are_retained(local_paths):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text = '组合形状中的正文。' * 20
    image = BytesIO()
    Image.new('RGB', (50, 50), 'red').save(image, format='PNG')
    image.seek(0)
    group.shapes.add_picture(image, Inches(1), Inches(3), Inches(1), Inches(1))
    path = local_paths.root / 'sample.pptx'
    deck.save(path)
    parsed = documents.parse_document(path, 'pptx')
    assert '组合形状中的正文' in parsed.blocks[0].text
    assert parsed.blocks[0].visual_needed


def test_pdf_shared_unused_images_do_not_trigger_vision(local_paths):
    pdf = fitz.open()
    for _ in range(2):
        page = pdf.new_page()
        page.insert_text((60, 100), 'Clearly readable source evidence. ' * 4)
    image = BytesIO()
    Image.new('RGB', (40, 40), 'red').save(image, format='PNG')
    pdf[0].insert_image(fitz.Rect(20, 150, 70, 200), stream=image.getvalue())
    # LibreOffice shares one resource dictionary across text and image pages.
    key, resources = pdf.xref_get_key(pdf[0].xref, 'Resources')
    pdf.xref_set_key(pdf[1].xref, 'Resources', resources)
    path = local_paths.root / 'sample.pdf'
    pdf.save(path)
    pdf.close()
    parsed = documents.parse_document(path, 'pdf')
    assert parsed.blocks[0].visual_needed
    assert not parsed.blocks[1].visual_needed


def test_pdf_background_rules_skip_but_charts_and_scans_remain(local_paths):
    pdf = fitz.open()
    for index in range(3):
        page = pdf.new_page()
        if index != 2:
            page.insert_text((60, 100), 'Source evidence without missing words. ' * 4)
        page.draw_rect(page.rect, fill=(1, 1, 1), overlay=False)
        page.draw_line((20, 30), (580, 30))
        if index == 1:
            page.draw_rect(fitz.Rect(150, 150, 220, 300), fill=(1, 0, 0))
    path = local_paths.root / 'sample.pdf'
    pdf.save(path)
    pdf.close()
    parsed = documents.parse_document(path, 'pdf')
    assert [block.visual_needed for block in parsed.blocks] == [False, True, True]


@pytest.mark.asyncio
async def test_ocr_engine_is_reused_outside_event_loop(monkeypatch):
    import rapidocr
    calls = []
    main_thread = threading.get_ident()
    class Engine:
        def __init__(self):
            calls.append(('init', threading.get_ident()))
        def __call__(self, path):
            calls.append(('run', threading.get_ident()))
            return SimpleNamespace(txts=['recognized'])
    monkeypatch.setattr(rapidocr, 'RapidOCR', Engine)
    monkeypatch.setattr(services, '_OCR_ENGINE', None)
    for _ in range(2):
        await services.asyncio.to_thread(services._read_ocr, 'unused.png')
    assert [kind for kind, _ in calls] == ['init', 'run', 'run']
    assert all(thread != main_thread for _, thread in calls)


@pytest.mark.asyncio
async def test_ingest_skips_repeated_auth_failure_and_finishes_index(local_paths, monkeypatch):
    db = Database(local_paths.root / 'test.db')
    db.initialize()
    db.execute("INSERT INTO notebooks(id,title,created_at,updated_at) VALUES('n','Fixture','now','now')")
    db.execute("INSERT INTO sources(id,notebook_id,revision_id,filename,media_type,size_bytes,sha256,blob_path,state,created_at,updated_at) VALUES('s','n','r','test.pdf','application/pdf',1,'h','unused','queued','now','now')")
    monkeypatch.setattr(services, 'DB', db)
    Image.new('RGB', (20, 20)).save(local_paths.root / 'image.png')
    parsed = documents.ParsedDocument(blocks=[documents.ParsedBlock('', {'page': i}, 'image.png', True) for i in range(3)])
    monkeypatch.setattr(services, 'parse_document', lambda *args: parsed)
    monkeypatch.setattr(services, 'provider_by_id', lambda key: {'capabilities': {'vision': True}})
    attempts = []
    async def describe(*args):
        attempts.append(1)
        raise services.ProviderError('Unauthorized', status=401)
    monkeypatch.setattr(services, 'describe_image', describe)
    monkeypatch.setattr(services, '_read_ocr', lambda path: SimpleNamespace(txts=['本页可核对的正文及数值 42。']))
    def encode(values):
        assert db.fetchone("SELECT state FROM sources WHERE id='s'")['state'] == 'processing'
        return [[1.] for _ in values]
    monkeypatch.setattr(services, 'EMBEDDINGS', SimpleNamespace(encode=encode, mode='test'))
    await services.ingest_source('s', image_policy={'mode': 'process', 'processors': ['vlm', 'ocr']}, image_provider_ids={'vlm': 'test'})
    source = db.fetchone("SELECT * FROM sources WHERE id='s'")
    assert source['state'] == 'ready' and len(attempts) == 1
    metadata = json_load(source['metadata_json'], {})
    assert metadata['vision_pages'] == 3 and metadata['chunk_count'] == 3
    assert metadata['ingest_timings']['total_seconds'] >= metadata['ingest_timings']['index_seconds']


def test_chart_cache_requires_complete_category_value_pairs():
    xml = b'''<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart><c:plotArea><c:barChart><c:ser>
    <c:tx><c:v>Change</c:v></c:tx><c:cat><c:strRef><c:f>Assets!A2:A3</c:f><c:strCache><c:ptCount val="2"/><c:pt idx="0"><c:v>Gold</c:v></c:pt><c:pt idx="1"><c:v>Coin</c:v></c:pt></c:strCache></c:strRef></c:cat>
    <c:val><c:numRef><c:f>Assets!B2:B3</c:f><c:numCache><c:ptCount val="2"/><c:pt idx="0"><c:v>-0.032</c:v></c:pt><c:pt idx="1"><c:v>0.024</c:v></c:pt></c:numCache></c:numRef></c:val>
    </c:ser></c:barChart></c:plotArea></c:chart></c:chartSpace>'''
    block = documents._cached_chart_block(xml, 'chart1.xml')
    assert block.locator['sheet'] == 'Assets'
    assert 'Gold | Change | -0.032' in block.text
    assert 'Coin | Change | 0.024' in block.text
    assert documents._cached_chart_block(xml.replace(b'<c:v>0.024</c:v>', b'<c:v></c:v>'), 'chart1.xml') is None
    assert documents._cached_chart_block(xml.replace(b'</c:ser>', b'<c:trendline/></c:ser>', 1), 'chart1.xml') is None


def test_substantive_chinese_paragraph_is_not_filtered_as_one_word():
    text = '资料说明指标变化必须结合统计区间进行判断，同时保留原文对于样本范围和适用条件的限制。' * 4
    assert retrieval.is_quality_chunk({'content': text, 'locator': {'kind': 'section', 'section': '指标解释'}})
    assert not retrieval.is_quality_chunk({'content': '目录\n' + text, 'locator': {}})


def test_invalid_office_preview_does_not_discard_body(local_paths, monkeypatch):
    doc = Document()
    doc.add_paragraph('正文仍然可以提取。')
    path = local_paths.root / 'test.docx'
    doc.save(path)
    invalid = local_paths.root / 'invalid.pdf'
    invalid.write_bytes(b'not a PDF')
    monkeypatch.setattr(documents, '_convert_office_to_pdf', lambda *args: invalid)
    parsed = documents.parse_document(path, 'invalid-preview')
    assert parsed.blocks[0].text == '正文仍然可以提取。'
    assert parsed.preview_path is None
    assert parsed.metadata['warnings'][0]['code'] == 'office_preview_unavailable'
